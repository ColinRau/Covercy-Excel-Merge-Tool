import streamlit as st
import zipfile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import io
import difflib
from typing import Any, Dict, List, Tuple


def _safe_load_workbook(file_like):
    """Attempt to load an XLSX; on parse failure, show guidance and return None."""
    try:
        return load_workbook(file_like, data_only=False)
    except (ValueError, InvalidFileException, zipfile.BadZipFile) as e:
        st.error("⚠️  This template can’t be read as-is. Please open it in Excel, choose “Save As…”, and upload the saved copy.")
        st.caption(f"Details: {e}")
        return None 


def patch_xlsx_workbook_xml_from_template(template_xlsx_bytes: bytes, generated_xlsx_bytes: bytes) -> bytes:
    """
    Covercy appears to validate workbook-level metadata/extensions in `xl/workbook.xml`.
    `openpyxl` can drop unknown extensions on save, causing Covercy to reject the file.

    This function replaces `xl/workbook.xml` in the generated XLSX with the one from
    the original Covercy template, while keeping all worksheet data from the generated file.
    """
    part_name = "xl/workbook.xml"

    with zipfile.ZipFile(io.BytesIO(template_xlsx_bytes), "r") as tzip:
        template_part = tzip.read(part_name)

    with zipfile.ZipFile(io.BytesIO(generated_xlsx_bytes), "r") as gzip:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, "w") as outzip:
            for info in gzip.infolist():
                data = gzip.read(info.filename)
                if info.filename == part_name:
                    data = template_part
                outzip.writestr(info, data)

    return out_buf.getvalue()


def normalize_entity_name_basic(value: Any) -> str:
    """
    Basic normalization for entity names:
    - convert to string
    - trim
    - collapse repeated whitespace
    - case-insensitive compare via lowercase
    """
    s = str(value or "").strip()
    s = " ".join(s.split())
    return s.lower()


def build_normalized_entity_index(target_entities: List[str]) -> Dict[str, List[str]]:
    """
    Build a lookup from normalized entity name -> list of original target entity strings.
    Multiple originals can share the same normalized key (ambiguous).
    """
    idx: Dict[str, List[str]] = {}
    for t in target_entities:
        key = normalize_entity_name_basic(t)
        if key == "":
            continue
        idx.setdefault(key, []).append(t)
    return idx


def suggest_target_entity(
    source_entity: Any,
    *,
    target_entities: List[str],
    target_index: Dict[str, List[str]],
    cutoff: float = 0.6,
) -> Tuple[str, str]:
    """
    Suggest a target entity for a given source entity.

    Strategy:
    - Exact match after basic normalization (if unique) -> ("<target>", "Exact")
    - Exact match after basic normalization (if ambiguous) -> ("", "Ambiguous")
    - Fallback to difflib fuzzy suggestion (current behavior) -> ("<target>", "Fuzzy")
    - No suggestion -> ("", "None")
    """
    src_raw = str(source_entity or "").strip()
    if src_raw == "":
        return "", "None"

    norm = normalize_entity_name_basic(src_raw)
    if norm != "" and norm in target_index:
        candidates = target_index[norm]
        if len(candidates) == 1:
            return candidates[0], "Exact"
        return "", "Ambiguous"

    fuzzy = (difflib.get_close_matches(src_raw, target_entities, n=1, cutoff=cutoff) or [""])[0]
    if fuzzy:
        return fuzzy, "Fuzzy"
    return "", "None"
